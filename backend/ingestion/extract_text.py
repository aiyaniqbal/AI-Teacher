"""Extract raw text from uploaded educational material (PDF, DOCX, PPTX, TXT)."""
from pathlib import Path

from pypdf import PdfReader
from docx import Document
from pptx import Presentation


def extract_pdf(path: str) -> str:
    reader = PdfReader(path)
    pages = []
    for page in reader.pages:
        text = page.extract_text() or ""
        pages.append(text)
    return "\n".join(pages)


def extract_docx(path: str) -> str:
    doc = Document(path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    return "\n".join(paragraphs)


def extract_pptx(path: str) -> str:
    prs = Presentation(path)
    slides_text = []
    for i, slide in enumerate(prs.slides, start=1):
        slide_lines = [f"[Slide {i}]"]
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in para.runs)
                    if line.strip():
                        slide_lines.append(line)
        slides_text.append("\n".join(slide_lines))
    return "\n\n".join(slides_text)


def extract_txt(path: str) -> str:
    return Path(path).read_text(encoding="utf-8", errors="ignore")


EXTRACTORS = {
    ".pdf": extract_pdf,
    ".docx": extract_docx,
    ".pptx": extract_pptx,
    ".txt": extract_txt,
}


def extract_text(path: str) -> str:
    """Dispatch to the correct extractor based on file extension."""
    ext = Path(path).suffix.lower()
    if ext not in EXTRACTORS:
        raise ValueError(f"Unsupported file type: {ext}")
    text = EXTRACTORS[ext](path)
    if not text.strip():
        raise ValueError(
            "No extractable text found in file (it may be a scanned/image PDF - "
            "add OCR later if you need to support those)."
        )
    return text
